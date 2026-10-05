import asyncio
from pathlib import Path

import httpx
import pytest
from openpyxl import load_workbook

from pricemon import report, storage
from pricemon.fetch import crawl, get
from pricemon.parse import Product, parse_catalogue, parse_price
from pricemon.site import BOOKS_TO_SCRAPE

PAGE = (Path(__file__).parent / "fixtures" / "catalogue_page.html").read_text(encoding="utf-8")
URL = "https://books.toscrape.com/catalogue/page-1.html"


@pytest.mark.parametrize(
    "text, expected",
    [("£51.77", 51.77), ("1 234,50 ₽", 1234.5), ("$1,299", 1299.0), ("1,299.99", 1299.99), ("990 руб.", 990.0)],
)
def test_parse_price(text, expected):
    assert parse_price(text) == expected


def test_parse_real_catalogue_page():
    products, next_url = parse_catalogue(PAGE, URL, BOOKS_TO_SCRAPE)
    assert len(products) == 20
    first = products[0]
    assert first.title == "A Light in the Attic"
    assert first.price == 51.77
    assert first.in_stock
    assert first.url == "https://books.toscrape.com/catalogue/a-light-in-the-attic_1000/index.html"
    assert next_url == "https://books.toscrape.com/catalogue/page-2.html"


def test_cards_without_price_are_skipped():
    html = '<article class="product_pod"><h3><a href="x" title="Реклама">Реклама</a></h3></article>'
    assert parse_catalogue(html, URL, BOOKS_TO_SCRAPE) == ([], None)


def product(name, price, in_stock=True):
    return Product(f"https://shop.test/{name}", name, price, in_stock)


def test_changes_between_runs(tmp_path):
    conn = storage.connect(str(tmp_path / "prices.db"))
    first = storage.save_run(conn, [product("a", 100), product("b", 200), product("c", 300), product("d", 50)])
    assert storage.changes(conn, first) == []
    second = storage.save_run(
        conn, [product("a", 120), product("b", 200, in_stock=False), product("d", 50), product("e", 10)]
    )
    found = {change.title: change for change in storage.changes(conn, second)}
    assert set(found) == {"a", "b", "c", "e"}
    assert (found["a"].kind, found["a"].percent) == ("price", 20.0)
    assert (found["b"].kind, found["b"].in_stock) == ("stock", False)
    assert found["c"].kind == "removed"
    assert found["e"].kind == "new"


def test_duplicate_product_in_one_run_is_stored_once(tmp_path):
    conn = storage.connect(str(tmp_path / "prices.db"))
    run_id = storage.save_run(conn, [product("a", 100), product("a", 100)])
    assert len(storage.load_run(conn, run_id)) == 1


def test_excel_report(tmp_path):
    conn = storage.connect(str(tmp_path / "prices.db"))
    storage.save_run(conn, [product("b", 200), product("a", 100)])
    run_id = storage.save_run(conn, [product("b", 150), product("a", 100)])
    found = storage.changes(conn, run_id)
    path = tmp_path / "report.xlsx"
    report.export_xlsx(conn, run_id, found, str(path))
    book = load_workbook(path)
    assert [row[0].value for row in book["Цены"].iter_rows(min_row=2)] == ["a", "b"]
    assert [cell.value for cell in book["Изменения"][2]][:5] == ["Изменилась цена", "b", 200, 150, -25]
    assert "b: 200 → 150 (-25%)" in report.summary(found)


def test_crawl_follows_pages_and_respects_limit():
    requested = []

    def handler(request):
        requested.append(str(request.url))
        return httpx.Response(200, text=PAGE)

    async def run(pages):
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            return await crawl(BOOKS_TO_SCRAPE, pages, delay=0, client=client)

    assert len(asyncio.run(run(2))) == 40
    assert requested == [URL, "https://books.toscrape.com/catalogue/page-2.html"]


def test_get_retries_then_fails(monkeypatch):
    calls = []

    async def no_sleep(_):
        pass

    monkeypatch.setattr(asyncio, "sleep", no_sleep)

    def handler(request):
        calls.append(1)
        return httpx.Response(503)

    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            await get(client, URL, attempts=3)

    with pytest.raises(RuntimeError, match="после 3 попыток"):
        asyncio.run(run())
    assert len(calls) == 3
