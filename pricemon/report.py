"""Выгрузка в Excel и текстовая сводка изменений."""
import sqlite3

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from pricemon.storage import Change, load_run

KINDS = {"new": "Новый товар", "removed": "Пропал из каталога", "price": "Изменилась цена", "stock": "Изменилось наличие"}
HEADER_FILL = PatternFill("solid", fgColor="1F3A4D")
UP, DOWN = Font(color="C0392B"), Font(color="1E8449")


def write_sheet(sheet, header, rows, widths):
    sheet.append(header)
    for cell in sheet[1]:
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), HEADER_FILL
    for row in rows:
        sheet.append(row)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def export_xlsx(conn: sqlite3.Connection, run_id: int, found: list[Change], path: str) -> None:
    book = Workbook()
    prices = book.active
    prices.title = "Цены"
    rows = sorted(load_run(conn, run_id).values(), key=lambda row: row["title"].lower())
    write_sheet(
        prices,
        ["Товар", "Цена", "В наличии", "Ссылка"],
        [[row["title"], row["price"], "да" if row["in_stock"] else "нет", row["url"]] for row in rows],
        [60, 12, 12, 70],
    )
    sheet = book.create_sheet("Изменения")
    write_sheet(
        sheet,
        ["Что изменилось", "Товар", "Было", "Стало", "Изменение, %", "Ссылка"],
        [[KINDS[c.kind], c.title, c.old_price, c.new_price, c.percent, c.url] for c in found],
        [22, 60, 12, 12, 14, 70],
    )
    for row, change in zip(sheet.iter_rows(min_row=2), found):
        if change.percent is not None:
            row[4].font = UP if change.percent > 0 else DOWN
    book.save(path)


def summary(found: list[Change], limit: int = 15) -> str:
    """Короткий текст для консоли или Telegram."""
    if not found:
        return "Изменений нет."
    lines = [f"Изменений: {len(found)}"]
    for change in found[:limit]:
        if change.kind == "price":
            lines.append(f"• {change.title}: {change.old_price:g} → {change.new_price:g} ({change.percent:+g}%)")
        elif change.kind == "stock":
            lines.append(f"• {change.title}: {'снова в наличии' if change.in_stock else 'нет в наличии'}")
        else:
            lines.append(f"• {KINDS[change.kind]}: {change.title}")
    if len(found) > limit:
        lines.append(f"…и ещё {len(found) - limit}")
    return "\n".join(lines)
